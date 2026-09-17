import streamlit as st
import os
import json
import hashlib
import re
from pathlib import Path
from typing import List

import requests
import chromadb
import PyPDF2
import docx
import openpyxl
import chardet

from sentence_transformers import SentenceTransformer


# ============================================================
# CONFIGURACIÓN
# ============================================================

st.set_page_config(
    page_title="SimpleRAG",
    page_icon="💬",
    layout="wide"
)


# ============================================================
# CONFIGURACIÓN GENERAL
# ============================================================

OLLAMA_CLOUD_URL = "https://ollama.com"

# Modelo para generación de respuestas
DEFAULT_CHAT_MODEL = "gpt-oss:20b"

# Modelo LOCAL para embeddings
DEFAULT_EMBEDDING_MODEL = (
    "sentence-transformers/all-MiniLM-L6-v2"
)

DEFAULT_CHUNK_SIZE = 1000
DEFAULT_CHUNK_OVERLAP = 150
DEFAULT_NUM_CHUNKS = 5
DEFAULT_TEMPERATURE = 0.1
DEFAULT_MEMORY_SIZE = 3

MAX_FILE_SIZE_MB = 25
MAX_SUMMARY_CHARS = 12000


# ============================================================
# OLLAMA CLOUD
# ============================================================

def get_ollama_api_key():

    try:

        key = st.secrets.get(
            "OLLAMA_API_KEY"
        )

        if key:
            return key

    except Exception:
        pass

    return os.getenv(
        "OLLAMA_API_KEY"
    )


def get_ollama_headers():

    api_key = get_ollama_api_key()

    if not api_key:
        return None

    return {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json"
    }


def validate_ollama_key():

    key = get_ollama_api_key()

    if not key:

        return (
            False,
            "No se encontró OLLAMA_API_KEY."
        )

    return (
        True,
        "API Key encontrada."
    )


# ============================================================
# OLLAMA CHAT
# ============================================================

def ollama_chat(
    model: str,
    prompt: str,
    temperature: float = 0.1,
    stream: bool = False
):

    headers = get_ollama_headers()

    if not headers:

        st.error(
            "No se encontró OLLAMA_API_KEY "
            "en Streamlit Secrets."
        )

        return None

    url = (
        f"{OLLAMA_CLOUD_URL}/api/chat"
    )

    payload = {
        "model": model,
        "messages": [
            {
                "role": "user",
                "content": prompt
            }
        ],
        "temperature": temperature,
        "stream": stream
    }

    try:

        response = requests.post(
            url,
            headers=headers,
            json=payload,
            stream=stream,
            timeout=300
        )

        response.raise_for_status()

        return response

    except requests.exceptions.HTTPError as e:

        error_text = ""

        try:
            error_text = response.text[:2000]
        except Exception:
            pass

        st.error(
            f"Error HTTP de Ollama Cloud:\n\n"
            f"{e}\n\n"
            f"{error_text}"
        )

        return None

    except requests.exceptions.RequestException as e:

        st.error(
            f"Error conectando con Ollama Cloud:\n\n{e}"
        )

        return None


# ============================================================
# MODELO DE EMBEDDINGS
# ============================================================

@st.cache_resource(
    show_spinner="Cargando modelo de embeddings..."
)
def load_embedding_model():

    return SentenceTransformer(
        DEFAULT_EMBEDDING_MODEL
    )


def embed_texts(
    texts: List[str]
):

    if not texts:
        return []

    try:

        model = load_embedding_model()

        embeddings = model.encode(
            texts,
            batch_size=32,
            show_progress_bar=False,
            normalize_embeddings=True
        )

        return embeddings.tolist()

    except Exception as e:

        st.error(
            f"Error generando embeddings: {e}"
        )

        return None


def embed_text(
    text: str
):

    embeddings = embed_texts(
        [text]
    )

    if not embeddings:

        return None

    return embeddings[0]


# ============================================================
# LECTURA DE PDF
# ============================================================

def read_pdf(file):

    text_parts = []

    try:

        reader = PyPDF2.PdfReader(
            file
        )

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:

                text_parts.append(
                    page_text
                )

        return "\n".join(
            text_parts
        )

    except Exception as e:

        st.error(
            f"Error leyendo PDF "
            f"{file.name}: {e}"
        )

        return ""


# ============================================================
# LECTURA DE WORD
# ============================================================

def read_docx(file):

    try:

        document = docx.Document(
            file
        )

        paragraphs = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                paragraphs.append(
                    text
                )

        return "\n".join(
            paragraphs
        )

    except Exception as e:

        st.error(
            f"Error leyendo DOCX "
            f"{file.name}: {e}"
        )

        return ""


# ============================================================
# LECTURA DE EXCEL
# ============================================================

def read_excel(file):

    try:

        workbook = openpyxl.load_workbook(
            file,
            read_only=True,
            data_only=True
        )

        text_parts = []

        for sheet in workbook.worksheets:

            text_parts.append(
                f"\n--- HOJA: {sheet.title} ---"
            )

            for row in sheet.iter_rows(
                values_only=True
            ):

                values = []

                for value in row:

                    if value is not None:

                        values.append(
                            str(value)
                        )

                if values:

                    text_parts.append(
                        " | ".join(values)
                    )

        return "\n".join(
            text_parts
        )

    except Exception as e:

        st.error(
            f"Error leyendo Excel "
            f"{file.name}: {e}"
        )

        return ""


# ============================================================
# LECTURA DE TXT / CSV
# ============================================================

def read_text_file(file):

    try:

        raw_data = file.read()

        detected = chardet.detect(
            raw_data
        )

        encoding = detected.get(
            "encoding",
            "utf-8"
        )

        return raw_data.decode(
            encoding,
            errors="replace"
        )

    except Exception as e:

        st.error(
            f"Error leyendo "
            f"{file.name}: {e}"
        )

        return ""


# ============================================================
# SELECTOR DE LECTOR
# ============================================================

def read_uploaded_file(file):

    extension = (
        Path(file.name)
        .suffix
        .lower()
    )

    if extension == ".pdf":

        return read_pdf(file)

    if extension == ".docx":

        return read_docx(file)

    if extension in [
        ".xlsx",
        ".xlsm",
        ".xltx"
    ]:

        return read_excel(file)

    if extension in [
        ".txt",
        ".csv",
        ".md"
    ]:

        return read_text_file(file)

    return read_text_file(file)


# ============================================================
# LIMPIEZA DE TEXTO
# ============================================================

def clean_text(text):

    if not text:

        return ""

    text = text.replace(
        "\x00",
        " "
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


# ============================================================
# CHUNKING
# ============================================================

def create_chunks(
    text,
    chunk_size,
    chunk_overlap
):

    text = clean_text(
        text
    )

    if not text:

        return []

    if chunk_overlap >= chunk_size:

        chunk_overlap = int(
            chunk_size * 0.15
        )

    chunks = []

    start = 0

    text_length = len(
        text
    )

    while start < text_length:

        end = min(
            start + chunk_size,
            text_length
        )

        chunk = text[
            start:end
        ].strip()

        if chunk:

            chunks.append(
                chunk
            )

        if end >= text_length:

            break

        start = (
            end -
            chunk_overlap
        )

    return chunks


# ============================================================
# HASH DE ARCHIVO
# ============================================================

def get_file_hash(file):

    file.seek(0)

    content = file.read()

    file.seek(0)

    return hashlib.md5(
        content
    ).hexdigest()


# ============================================================
# RESUMEN DEL DOCUMENTO
# ============================================================

def summarize_document(
    model,
    text
):

    if not text:

        return ""

    text_for_summary = (
        text[:MAX_SUMMARY_CHARS]
    )

    prompt = f"""
Eres un asistente especializado
en análisis documental.

Resume el siguiente documento
de manera concisa.

Incluye cuando estén disponibles:

- Tema principal
- Información importante
- Datos relevantes
- Procesos
- Fechas
- Personas o áreas involucradas
- Conclusiones

No inventes información.

DOCUMENTO:

{text_for_summary}
"""

    response = ollama_chat(
        model=model,
        prompt=prompt,
        temperature=0.0,
        stream=False
    )

    if response is None:

        return ""

    try:

        data = response.json()

        message = data.get(
            "message",
            {}
        )

        return message.get(
            "content",
            ""
        ).strip()

    except Exception:

        return ""


# ============================================================
# TOKENIZACIÓN
# ============================================================

def tokenize(text):

    return set(
        re.findall(
            r"\b\w+\b",
            text.lower()
        )
    )


# ============================================================
# KEYWORD SCORE
# ============================================================

def keyword_score(
    query,
    document
):

    query_words = tokenize(
        query
    )

    document_words = tokenize(
        document
    )

    if not query_words:

        return 0.0

    intersection = (
        query_words &
        document_words
    )

    return (
        len(intersection) /
        len(query_words)
    )


# ============================================================
# NORMALIZAR SCORES
# ============================================================

def normalize_scores(
    scores
):

    if not scores:

        return []

    minimum = min(
        scores
    )

    maximum = max(
        scores
    )

    if maximum == minimum:

        return [
            1.0 if maximum > 0
            else 0.0
            for _ in scores
        ]

    return [
        (
            score - minimum
        ) /
        (
            maximum - minimum
        )
        for score in scores
    ]


# ============================================================
# CREAR CHROMADB
# ============================================================

def create_vector_collection():

    client = chromadb.Client()

    collection_name = (
        "simplerag_collection"
    )

    try:

        client.delete_collection(
            collection_name
        )

    except Exception:

        pass

    collection = client.create_collection(
        name=collection_name
    )

    return collection


# ============================================================
# PROCESAR DOCUMENTOS
# ============================================================

def process_documents(
    uploaded_files,
    chat_model,
    chunk_size,
    chunk_overlap
):

    collection = (
        create_vector_collection()
    )

    all_chunks = []
    all_metadata = []
    all_ids = []

    summaries = {}
    documents_info = []

    progress = st.progress(
        0
    )

    total_files = len(
        uploaded_files
    )

    for file_index, file in enumerate(
        uploaded_files
    ):

        # ----------------------------------------------------
        # Validar tamaño
        # ----------------------------------------------------

        file.seek(0)

        file_bytes = file.read()

        file.seek(0)

        file_size_mb = (
            len(file_bytes) /
            1024 /
            1024
        )

        if file_size_mb > MAX_FILE_SIZE_MB:

            st.warning(
                f"{file.name} supera "
                f"{MAX_FILE_SIZE_MB} MB "
                f"y fue omitido."
            )

            continue

        # ----------------------------------------------------
        # Hash
        # ----------------------------------------------------

        file_hash = get_file_hash(
            file
        )

        # ----------------------------------------------------
        # Leer documento
        # ----------------------------------------------------

        with st.spinner(
            f"Leyendo {file.name}..."
        ):

            text = read_uploaded_file(
                file
            )

        text = clean_text(
            text
        )

        if not text:

            st.warning(
                f"No se encontró texto "
                f"en {file.name}."
            )

            continue

        # ----------------------------------------------------
        # Crear resumen
        # ----------------------------------------------------

        with st.spinner(
            f"Analizando {file.name}..."
        ):

            summary = summarize_document(
                chat_model,
                text
            )

        summaries[
            file.name
        ] = summary

        # ----------------------------------------------------
        # Crear chunks
        # ----------------------------------------------------

        chunks = create_chunks(
            text,
            chunk_size,
            chunk_overlap
        )

        documents_info.append({

            "name": file.name,

            "hash": file_hash,

            "characters": len(
                text
            ),

            "chunks": len(
                chunks
            )
        })

        # ----------------------------------------------------
        # Agregar chunks
        # ----------------------------------------------------

        for chunk_index, chunk in enumerate(
            chunks
        ):

            chunk_id = (
                f"{file_hash}_"
                f"{chunk_index}"
            )

            all_chunks.append(
                chunk
            )

            all_metadata.append({

                "source":
                    file.name,

                "chunk":
                    chunk_index
            })

            all_ids.append(
                chunk_id
            )

        progress.progress(
            (
                file_index + 1
            ) /
            total_files
        )

    progress.empty()

    if not all_chunks:

        st.error(
            "No se pudieron procesar "
            "los documentos."
        )

        return (
            None,
            {},
            []
        )

    # ========================================================
    # EMBEDDINGS LOCALES
    # ========================================================

    st.info(
        f"Generando embeddings locales "
        f"para {len(all_chunks)} chunks..."
    )

    embedding_progress = st.progress(
        0
    )

    embeddings = []

    batch_size = 32

    total_batches = (
        len(all_chunks) +
        batch_size -
        1
    ) // batch_size

    for batch_number, start in enumerate(
        range(
            0,
            len(all_chunks),
            batch_size
        )
    ):

        batch = all_chunks[
            start:
            start + batch_size
        ]

        batch_embeddings = embed_texts(
            batch
        )

        if not batch_embeddings:

            st.error(
                "No se pudieron generar "
                "los embeddings."
            )

            return (
                None,
                {},
                []
            )

        embeddings.extend(
            batch_embeddings
        )

        embedding_progress.progress(
            (
                batch_number + 1
            ) /
            total_batches
        )

    embedding_progress.empty()

    # ========================================================
    # INSERTAR EN CHROMA
    # ========================================================

    collection.add(

        ids=all_ids,

        embeddings=embeddings,

        documents=all_chunks,

        metadatas=all_metadata
    )

    return (
        collection,
        summaries,
        documents_info
    )


# ============================================================
# BÚSQUEDA HÍBRIDA
# ============================================================

def hybrid_search(
    collection,
    query,
    num_chunks
):

    if collection is None:

        return []

    # --------------------------------------------------------
    # Embedding de la pregunta
    # --------------------------------------------------------

    query_embedding = embed_text(
        query
    )

    if not query_embedding:

        return []

    # --------------------------------------------------------
    # Cantidad de candidatos
    # --------------------------------------------------------

    collection_count = (
        collection.count()
    )

    candidate_count = min(
        max(
            num_chunks * 4,
            10
        ),
        collection_count
    )

    # --------------------------------------------------------
    # Búsqueda vectorial
    # --------------------------------------------------------

    results = collection.query(

        query_embeddings=[
            query_embedding
        ],

        n_results=candidate_count,

        include=[
            "documents",
            "metadatas",
            "distances"
        ]
    )

    documents = (
        results
        .get(
            "documents",
            [[]]
        )[0]
    )

    metadatas = (
        results
        .get(
            "metadatas",
            [[]]
        )[0]
    )

    distances = (
        results
        .get(
            "distances",
            [[]]
        )[0]
    )

    if not documents:

        return []

    # --------------------------------------------------------
    # Semantic score
    # --------------------------------------------------------

    semantic_scores = [

        1 /
        (
            1 + distance
        )

        for distance in distances
    ]

    # --------------------------------------------------------
    # Keyword score
    # --------------------------------------------------------

    keyword_scores = [

        keyword_score(
            query,
            document
        )

        for document in documents
    ]

    semantic_normalized = (
        normalize_scores(
            semantic_scores
        )
    )

    keyword_normalized = (
        normalize_scores(
            keyword_scores
        )
    )

    # --------------------------------------------------------
    # Score híbrido
    # --------------------------------------------------------

    combined = []

    for i in range(
        len(documents)
    ):

        score = (

            0.70 *
            semantic_normalized[i]

            +

            0.30 *
            keyword_normalized[i]
        )

        combined.append({

            "document":
                documents[i],

            "metadata":
                metadatas[i],

            "semantic_score":
                semantic_normalized[i],

            "keyword_score":
                keyword_normalized[i],

            "score":
                score
        })

    combined.sort(
        key=lambda x:
            x["score"],
        reverse=True
    )

    return combined[
        :num_chunks
    ]


# ============================================================
# CONSTRUIR PROMPT RAG
# ============================================================

def build_rag_prompt(
    question,
    retrieved_chunks,
    summaries,
    chat_history,
    prompt_template
):

    context_parts = []

    for index, item in enumerate(
        retrieved_chunks,
        start=1
    ):

        source = (
            item[
                "metadata"
            ]
            .get(
                "source",
                "Documento"
            )
        )

        context_parts.append(
            f"""
--- FUENTE {index}: {source} ---

{item["document"]}
"""
        )

    context = "\n".join(
        context_parts
    )

    # --------------------------------------------------------
    # Resúmenes
    # --------------------------------------------------------

    summary_text = ""

    for filename, summary in (
        summaries.items()
    ):

        if summary:

            summary_text += (
                f"\n--- {filename} ---\n"
                f"{summary}\n"
            )

    # --------------------------------------------------------
    # Historial
    # --------------------------------------------------------

    history_text = ""

    for message in (
        chat_history[-6:]
    ):

        role = message.get(
            "role",
            ""
        )

        content = message.get(
            "content",
            ""
        )

        history_text += (
            f"{role.upper()}: "
            f"{content}\n"
        )

    # --------------------------------------------------------
    # Prompt
    # --------------------------------------------------------

    prompt = f"""
{prompt_template}

========================
CONTEXTO RECUPERADO
========================

{context}

========================
RESÚMENES
========================

{summary_text}

========================
HISTORIAL
========================

{history_text}

========================
PREGUNTA
========================

{question}

========================
RESPUESTA
========================
"""

    return prompt


# ============================================================
# STREAM DE RESPUESTA OLLAMA
# ============================================================

def stream_ollama_response(
    response
):

    for line in response.iter_lines():

        if not line:

            continue

        try:

            decoded = line.decode(
                "utf-8"
            )

            data = json.loads(
                decoded
            )

            message = data.get(
                "message",
                {}
            )

            content = message.get(
                "content",
                ""
            )

            if content:

                yield content

        except Exception:

            continue


# ============================================================
# SESSION STATE
# ============================================================

def initialize_session_state():

    defaults = {

        "messages": [],

        "collection": None,

        "summaries": {},

        "documents_info": [],

        "data_processed": False,

        "chat_model":
            DEFAULT_CHAT_MODEL
    }

    for key, value in (
        defaults.items()
    ):

        if key not in st.session_state:

            st.session_state[
                key
            ] = value


initialize_session_state()


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:

    st.header(
        "⚙️ Configuración"
    )

    # ========================================================
    # OLLAMA CLOUD
    # ========================================================

    st.subheader(
        "☁️ Ollama Cloud"
    )

    key_ok, key_message = (
        validate_ollama_key()
    )

    if key_ok:

        st.success(
            "OLLAMA_API_KEY detectada"
        )

    else:

        st.error(
            key_message
        )

    st.caption(
        "Ollama Cloud se utiliza para "
        "generar las respuestas."
    )

    # ========================================================
    # MODELO CHAT
    # ========================================================

    st.subheader(
        "🤖 Modelo de Chat"
    )

    chat_model = st.text_input(

        "Modelo",

        value=(
            st.session_state.chat_model
        ),

        help=(
            "Modelo utilizado por "
            "Ollama Cloud."
        )
    )

    st.session_state.chat_model = (
        chat_model
    )

    # ========================================================
    # EMBEDDINGS
    # ========================================================

    st.subheader(
        "🧠 Embeddings"
    )

    st.info(
        "Embeddings locales:\n\n"
        "all-MiniLM-L6-v2"
    )

    # ========================================================
    # CONFIGURACIÓN RAG
    # ========================================================

    st.subheader(
        "📚 Configuración RAG"
    )

    chunk_size = st.slider(

        "Tamaño del chunk",

        min_value=500,

        max_value=2000,

        value=DEFAULT_CHUNK_SIZE,

        step=100
    )

    chunk_overlap = st.slider(

        "Overlap",

        min_value=50,

        max_value=400,

        value=DEFAULT_CHUNK_OVERLAP,

        step=50
    )

    num_chunks = st.slider(

        "Chunks recuperados",

        min_value=2,

        max_value=10,

        value=DEFAULT_NUM_CHUNKS
    )

    temperature = st.slider(

        "Temperature",

        min_value=0.0,

        max_value=1.0,

        value=DEFAULT_TEMPERATURE,

        step=0.05
    )

    memory_size = st.slider(

        "Memoria",

        min_value=1,

        max_value=10,

        value=DEFAULT_MEMORY_SIZE
    )

    # ========================================================
    # PROMPT
    # ========================================================

    st.subheader(
        "📝 Prompt"
    )

    prompt_template = st.text_area(

        "Instrucciones",

        value="""
Responde la pregunta utilizando
exclusivamente la información encontrada
en los documentos.

Si la información no está disponible,
indica claramente que no se encuentra
en los documentos proporcionados.

No inventes datos.

Responde en español de manera clara,
profesional y concisa.
""",

        height=180
    )

    # ========================================================
    # DOCUMENTOS
    # ========================================================

    st.subheader(
        "📂 Documentos"
    )

    uploaded_files = st.file_uploader(

        "Sube tus documentos",

        type=[
            "pdf",
            "docx",
            "xlsx",
            "xlsm",
            "txt",
            "csv",
            "md"
        ],

        accept_multiple_files=True
    )

    process_button = st.button(

        "🚀 Procesar documentos",

        type="primary",

        use_container_width=True
    )

    # ========================================================
    # LIMPIAR CHAT
    # ========================================================

    if st.button(

        "🗑️ Limpiar conversación",

        use_container_width=True
    ):

        st.session_state.messages = []

        st.rerun()


# ============================================================
# PROCESAR DOCUMENTOS
# ============================================================

if process_button:

    if not key_ok:

        st.error(
            "Configura OLLAMA_API_KEY "
            "en Streamlit Secrets."
        )

        st.stop()

    if not uploaded_files:

        st.warning(
            "Sube al menos un documento."
        )

        st.stop()

    with st.spinner(
        "Procesando documentos..."
    ):

        (
            collection,
            summaries,
            documents_info
        ) = process_documents(

            uploaded_files=
                uploaded_files,

            chat_model=
                chat_model,

            chunk_size=
                chunk_size,

            chunk_overlap=
                chunk_overlap
        )

    if collection is not None:

        st.session_state.collection = (
            collection
        )

        st.session_state.summaries = (
            summaries
        )

        st.session_state.documents_info = (
            documents_info
        )

        st.session_state.data_processed = (
            True
        )

        st.success(
            "✅ Documentos procesados correctamente."
        )


# ============================================================
# TÍTULO
# ============================================================

st.title(
    "💬 SimpleRAG"
)

st.caption(
    "RAG documental con "
    "Streamlit + Ollama Cloud + "
    "Sentence Transformers + ChromaDB"
)


# ============================================================
# MÉTRICAS
# ============================================================

if st.session_state.data_processed:

    total_chunks = (

        st.session_state
        .collection
        .count()

        if st.session_state.collection

        else 0
    )

    col1, col2, col3 = (
        st.columns(3)
    )

    with col1:

        st.metric(

            "Documentos",

            len(
                st.session_state
                .documents_info
            )
        )

    with col2:

        st.metric(
            "Chunks",
            total_chunks
        )

    with col3:

        st.metric(
            "Modelo LLM",
            chat_model
        )


# ============================================================
# DOCUMENTOS
# ============================================================

if st.session_state.documents_info:

    with st.expander(
        "📚 Ver documentos procesados"
    ):

        for document in (
            st.session_state
            .documents_info
        ):

            st.write(
                f"**{document['name']}**"
            )

            st.caption(

                f"Caracteres: "
                f"{document['characters']:,} | "

                f"Chunks: "
                f"{document['chunks']}"
            )


# ============================================================
# CHAT
# ============================================================

if not st.session_state.data_processed:

    st.info(
        "👈 Sube documentos desde "
        "el panel lateral y presiona "
        "**Procesar documentos**."
    )

else:

    st.subheader(
        "💬 Pregunta sobre tus documentos"
    )

    # --------------------------------------------------------
    # Mostrar mensajes
    # --------------------------------------------------------

    for message in (
        st.session_state.messages
    ):

        with st.chat_message(
            message["role"]
        ):

            st.markdown(
                message["content"]
            )

    # --------------------------------------------------------
    # Input
    # --------------------------------------------------------

    question = st.chat_input(
        "Escribe tu pregunta..."
    )

    if question:

        # ====================================================
        # MENSAJE USUARIO
        # ====================================================

        st.session_state.messages.append({

            "role":
                "user",

            "content":
                question
        })

        with st.chat_message(
            "user"
        ):

            st.markdown(
                question
            )

        # ====================================================
        # RETRIEVAL
        # ====================================================

        with st.spinner(
            "🔎 Buscando información relevante..."
        ):

            retrieved_chunks = (
                hybrid_search(

                    collection=
                        st.session_state.collection,

                    query=
                        question,

                    num_chunks=
                        num_chunks
                )
            )

        if not retrieved_chunks:

            answer = (
                "No encontré información "
                "relevante en los documentos."
            )

            with st.chat_message(
                "assistant"
            ):

                st.markdown(
                    answer
                )

            st.session_state.messages.append({

                "role":
                    "assistant",

                "content":
                    answer
            })

            st.stop()

        # ====================================================
        # CONSTRUIR PROMPT
        # ====================================================

        prompt = build_rag_prompt(

            question=
                question,

            retrieved_chunks=
                retrieved_chunks,

            summaries=
                st.session_state.summaries,

            chat_history=
                st.session_state.messages[
                    -(memory_size * 2):
                ],

            prompt_template=
                prompt_template
        )

        # ====================================================
        # OLLAMA CLOUD
        # ====================================================

        with st.chat_message(
            "assistant"
        ):

            response = ollama_chat(

                model=
                    chat_model,

                prompt=
                    prompt,

                temperature=
                    temperature,

                stream=
                    True
            )

            if response is None:

                st.stop()

            answer_placeholder = (
                st.empty()
            )

            answer = ""

            for token in (
                stream_ollama_response(
                    response
                )
            ):

                answer += token

                answer_placeholder.markdown(
                    answer
                )

        # ====================================================
        # GUARDAR RESPUESTA
        # ====================================================

        st.session_state.messages.append({

            "role":
                "assistant",

            "content":
                answer
        })

        # ====================================================
        # FUENTES
        # ====================================================

        with st.expander(
            "🔎 Ver información recuperada"
        ):

            for index, item in enumerate(

                retrieved_chunks,

                start=1
            ):

                source = (

                    item[
                        "metadata"
                    ]
                    .get(
                        "source",
                        "Documento"
                    )
                )

                score = item[
                    "score"
                ]

                st.markdown(
                    f"### {index}. {source}"
                )

                st.caption(
                    f"Relevancia: "
                    f"{score:.2%}"
                )

                st.write(
                    item["document"]
                )

                st.divider()


# ============================================================
# FOOTER
# ============================================================

st.divider()

st.caption(
    "SimpleRAG • Streamlit Cloud • "
    "Ollama Cloud • Sentence Transformers • "
    "ChromaDB"
)
